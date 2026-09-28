from flask import Flask, render_template, request, redirect, url_for, send_from_directory, session
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError
import os
from werkzeug.utils import secure_filename
from functools import wraps
from threading import Lock
from waitress import serve

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY")
OWNER_PASSWORD = os.environ.get("OWNER_PASSWORD")
DATA_LOCK = Lock()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "participants.xlsx")
PHOTO_FOLDER = os.path.join(BASE_DIR, "photos")
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

# Ensure photos folder exists
os.makedirs(PHOTO_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = PHOTO_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def save_photo_as_jpeg(photo, filename):
    """Decode the upload and save it as a standard RGB JPEG."""
    photo_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    try:
        with Image.open(photo.stream) as image:
            if image.mode in ('RGBA', 'LA'):
                background = Image.new('RGB', image.size, 'white')
                background.paste(image, mask=image.getchannel('A'))
                image = background
            else:
                image = image.convert('RGB')
            image.save(photo_path, format='JPEG', quality=92)
    except (UnidentifiedImageError, OSError) as error:
        raise ValueError("The uploaded file is not a valid image") from error
    return photo_path

def login_required(f):
    """Decorator to require owner login"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'owner' not in session:
            return redirect(url_for('owner_login'))
        return f(*args, **kwargs)
    return decorated_function

def delete_participant_photo(participant_number):
    """Delete any existing photo files for a participant"""
    photo_folder = app.config['UPLOAD_FOLDER']
    for filename in os.listdir(photo_folder):
        if filename.startswith(participant_number):
            filepath = os.path.join(photo_folder, filename)
            try:
                os.remove(filepath)
            except:
                pass

@app.route("/photos/<filename>")
def get_photo(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route("/owner/login", methods=["GET", "POST"])
def owner_login():
    if request.method == "POST":
        password = request.form.get("password")
        if password == OWNER_PASSWORD:
            session['owner'] = True
            return redirect(url_for('view_participants'))
        else:
            return render_template("owner_login.html", error="Invalid password")
    
    return render_template("owner_login.html")

@app.route("/owner/logout")
def owner_logout():
    session.pop('owner', None)
    return redirect(url_for('index'))

def create_excel_file():
    """Create the spreadsheet if it doesn't already exist."""
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Participants"

        sheet.append([
            "Participant Number",
            "Name",
            "Surname",
            "Actual Height (cm)",
            "Computer Measured Height (cm)",
            "Error (cm)",
            "Photo"
        ])

        # Set column widths
        sheet.column_dimensions['A'].width = 15
        sheet.column_dimensions['B'].width = 15
        sheet.column_dimensions['C'].width = 15
        sheet.column_dimensions['D'].width = 20
        sheet.column_dimensions['E'].width = 25
        sheet.column_dimensions['F'].width = 15
        sheet.column_dimensions['G'].width = 25

        workbook.save(EXCEL_FILE)
        workbook.close()

def get_next_participant_number():
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Participants"]

    # Minus the header row
    participant_count = sheet.max_row - 1

    next_number = participant_count + 1

    return f"P{next_number:03d}"

@app.route("/", methods=["GET", "POST"])
def index():

    if request.method == "POST":
        # Check if photo was uploaded
        if 'photo' not in request.files:
            return render_template("index.html", error="❌ No photo uploaded. Please choose a photo.")
        
        photo = request.files['photo']
        
        if photo.filename == '':
            return render_template("index.html", error="❌ No photo selected. Please choose a photo.")
        
        if not allowed_file(photo.filename):
            return render_template("index.html", error="❌ Invalid file type. Use PNG, JPG, JPEG, or GIF")

        name = request.form.get("name", "").strip()
        surname = request.form.get("surname", "").strip()
        
        if not name or not surname:
            return render_template("index.html", error="❌ Please enter both name and surname.")
        
        try:
            actual_height = float(request.form["actual_height"])
            if actual_height < 50 or actual_height > 250:
                return render_template("index.html", error="❌ Height must be between 50-250 cm.")
        except ValueError:
            return render_template("index.html", error="❌ Please enter a valid height in cm.")
        
        with DATA_LOCK:
            participant_number = get_next_participant_number()

            # Normalize every accepted upload to P00X.jpg.
            photo_filename = f"{participant_number}.jpg"
            try:
                photo_path = save_photo_as_jpeg(photo, photo_filename)
            except ValueError as error:
                return render_template("index.html", error=f"❌ {error}")

            workbook = load_workbook(EXCEL_FILE)
            sheet = workbook["Participants"]

            sheet.append([
                participant_number,
                name,
                surname,
                actual_height,
                "",  # Computer Measured Height (empty initially)
                "",  # Error (empty initially)
                photo_filename  # Store filename for reference
            ])

            # Add photo to Excel
            if os.path.exists(photo_path):
                try:
                    img = XLImage(photo_path)
                    img.width = 100
                    img.height = 100
                    # Add image to the photo column (G) of the new row
                    sheet.add_image(img, f'G{sheet.max_row}')
                    # Adjust row height to fit image
                    sheet.row_dimensions[sheet.max_row].height = 105
                except Exception as e:
                    print(f"Warning: Could not embed image: {e}")

            workbook.save(EXCEL_FILE)
            workbook.close()

        return redirect(
            url_for(
                "success",
                participant_number=participant_number
            )
        )

    return render_template("index.html")

@app.route("/success")
def success():
    participant_number = request.args.get("participant_number")
    
    # Calculate ranking stats
    create_excel_file()
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Participants"]
    
    # Get all participant heights
    all_heights = []
    current_height = None
    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:  # If participant number exists
            height = row[3]  # Actual height in column D
            if height:
                all_heights.append(height)
                if row[0] == participant_number:
                    current_height = height
    
    # Calculate ranking stats
    rank_data = {}
    if current_height and all_heights:
        # Total participants
        total = len(all_heights)
        
        # Count how many are taller
        taller_count = sum(1 for h in all_heights if h > current_height)
        
        # Calculate rank (1-based, where rank 1 is the tallest)
        rank = taller_count + 1
        
        # Calculate percentile (what % are shorter than this person)
        shorter_count = total - taller_count - 1
        percentile = (shorter_count / total * 100) if total > 1 else 0
        
        rank_data = {
            'rank': rank,
            'total': total,
            'percentile': round(percentile, 1),
            'height': current_height,
            'is_tallest': rank == 1,
            'is_shortest': taller_count == total - 1
        }
    
    workbook.close()
    return render_template("success.html", participant_number=participant_number, rank_data=rank_data)

@app.route("/participants")
@login_required
def view_participants():
    # Create Excel file if it doesn't exist
    create_excel_file()
    
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Participants"]
    
    participants = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:  # If participant number exists
            participants.append({
                'number': row[0],
                'name': row[1],
                'surname': row[2],
                'height': row[3],
                'photo': row[6]
            })
    
    workbook.close()
    return render_template("participants.html", participants=participants)

@app.route("/edit/<participant_number>", methods=["GET", "POST"])
@login_required
def edit_participant(participant_number):
    # Create Excel file if it doesn't exist
    create_excel_file()
    
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Participants"]
    
    # Find the participant row
    participant_row = None
    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
        if row[0].value == participant_number:
            participant_row = idx
            break
    
    if not participant_row:
        return {"error": "Participant not found"}, 404
    
    # Handle GET requests by redirecting to participants view
    if request.method == "GET":
        return redirect(url_for("view_participants"))
    
    name = request.form.get("name")
    surname = request.form.get("surname")
    actual_height = request.form.get("actual_height")
    computer_height = request.form.get("computer_height")
    
    # Update name, surname, height
    if name:
        sheet[f'B{participant_row}'].value = name
    if surname:
        sheet[f'C{participant_row}'].value = surname
    if actual_height:
        sheet[f'D{participant_row}'].value = float(actual_height)
    if computer_height:
        computer_height_float = float(computer_height)
        sheet[f'E{participant_row}'].value = computer_height_float
        # Calculate error: Computer Height - Actual Height
        actual_h = float(actual_height) if actual_height else float(sheet[f'D{participant_row}'].value)
        error = computer_height_float - actual_h
        sheet[f'F{participant_row}'].value = error
    
    # Handle photo upload
    if 'photo' in request.files:
        photo = request.files['photo']
        if photo and photo.filename != '' and allowed_file(photo.filename):
            # Delete old photos first to avoid conflicts
            delete_participant_photo(participant_number)
            
            # Save new photo with participant number
            file_ext = os.path.splitext(photo.filename)[1].lower()
            photo_filename = f"{participant_number}{file_ext}"
            photo_path = os.path.join(app.config['UPLOAD_FOLDER'], photo_filename)
            photo.save(photo_path)
            
            # Update photo filename in sheet
            sheet[f'G{participant_row}'].value = photo_filename
            
            # Update photo image in Excel
            if os.path.exists(photo_path):
                try:
                    img = XLImage(photo_path)
                    img.width = 100
                    img.height = 100
                    sheet.add_image(img, f'G{participant_row}')
                    sheet.row_dimensions[participant_row].height = 105
                except:
                    pass  # If image insertion fails, continue without it
    
    workbook.save(EXCEL_FILE)
    
    # Return updated participant data
    row = sheet[participant_row]
    result = {
        "success": True,
        "participant": {
            "number": row[0].value,
            "name": row[1].value,
            "surname": row[2].value,
            "height": row[3].value,
            "photo": row[6].value
        }
    }
    workbook.close()
    return result

@app.route("/delete/<participant_number>", methods=["POST"])
@login_required
def delete_participant(participant_number):
    # Create Excel file if it doesn't exist
    create_excel_file()
    
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Participants"]
    
    # Find and delete the participant row
    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
        if row[0].value == participant_number:
            # Delete the row
            sheet.delete_rows(idx, 1)
            workbook.save(EXCEL_FILE)
            workbook.close()
            return redirect(url_for("view_participants"))
    
    workbook.close()
    return redirect(url_for("view_participants"))

if __name__ == "__main__":
    missing_settings = [
        name for name, value in {
            "FLASK_SECRET_KEY": app.secret_key,
            "OWNER_PASSWORD": OWNER_PASSWORD,
        }.items() if not value
    ]
    if missing_settings:
        raise RuntimeError(f"Set required environment variables: {', '.join(missing_settings)}")

    create_excel_file()
    host = os.environ.get("WEB_HOST", "127.0.0.1")
    port = int(os.environ.get("WEB_PORT", "5000"))
    serve(app, host=host, port=port)
